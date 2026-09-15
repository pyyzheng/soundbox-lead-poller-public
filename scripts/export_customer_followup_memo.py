#!/usr/bin/env python3
"""按「最终分配的业务员」导出客户推进备忘录（Memo 模板版式）。

规则：
- 一行一客户（线索）
- 客户区：单位=Customer Name 客戶單位；对接人=Customer Name（客户名称）；业务员=最终分配
- 交流区：从 Follow-up Records 按时间升序展开；默认至少 2 次列块，≥3 次动态加列
- 模板结构对齐：Memo 客戶推進備忘記錄_中英文對照.xlsx

用法：
  python3 scripts/export_customer_followup_memo.py
  python3 scripts/export_customer_followup_memo.py --salesperson Gigi
  python3 scripts/export_customer_followup_memo.py --salesperson Gigi --out /tmp/memo_gigi.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

APP_TOKEN = os.environ.get("FEISHU_BITABLE_APP") or os.environ.get("FEISHU_APP_TOKEN") or "ZpbUb7SP7azsNasniFjc0bWSnHg"
CASE_TABLE = "tbluuuXn9WexH8LV"
FOLLOWUP_TABLE = "tbl3n8TTJYXHG12q"
TZ = ZoneInfo("Asia/Shanghai")

FIELD_LEAD_ID = "Clue ID"
FIELD_CUSTOMER = "Customer Name（客户名称）"
FIELD_UNIT = "Customer Name 客戶單位"
FIELD_FINAL = "The final assigned salesperson（最终分配的业务员）"
FIELD_ADDRESS = "单位地址 Address"
FIELD_EMAIL = "Email（客户邮箱）"
FIELD_PHONE = "Phone（客户电话）"
FIELD_WECHAT = "Wechat（微信）"
FIELD_CHANNEL = "Channels（渠道）"
FIELD_SUB = "Channel segmentation (细分渠道)"
FIELD_ENTRY = "Entry Time（录入时间）"
FIELD_CUST_TYPE = "Customer type"
FIELD_CUST_TYPE_US = "Customer Type（US）"

FU_TIME = "Follow-up Time"
FU_METHOD = "Contact Method"
FU_LOCATION = "地点 Location"
FU_ATTENDEES = "参会人 Attendees"
FU_PRODUCTS = "意向产品 Products"
FU_REGION = "合作区域 Region"
FU_QUOTE = "报价折扣 Quote/Disc"
FU_OTHER = "Follow-up Details"
FU_GIFTS = "赠送礼品 Gifts"
FU_NEXT = "Next Step"
FU_RELATED = "Related Lead"
FU_DELETED = "已删除"

COMM_HEADERS = [
    ("时间\nDate", FU_TIME),
    ("方式\nMethod", FU_METHOD),
    ("地点\nLocation", FU_LOCATION),
    ("参会人\nAttendees", FU_ATTENDEES),
    ("意向产品\nProducts", FU_PRODUCTS),
    ("合作区域\nRegion", FU_REGION),
    ("报价折扣\nQuote/Disc.", FU_QUOTE),
    ("其他事项\nOther", FU_OTHER),
    ("赠送礼品\nGifts", FU_GIFTS),
    ("推进计划\nNext Step", FU_NEXT),
]

CUSTOMER_HEADERS = [
    "序号\nNo.",
    "客户单位\nCustomer",
    "单位地址\nAddress",
    "对接人\nContact",
    "联系方式\nContact Info",
    "商机来源\nLead Source",
    "备案时间\nReg. Date",
    "客户类型\nCustomer Type",
]


def _token() -> str:
    app_id = os.environ["FEISHU_APP_ID"]
    app_secret = os.environ["FEISHU_APP_SECRET"]
    r = requests.post(
        "https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
        json={"app_id": app_id, "app_secret": app_secret},
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    if data.get("code") != 0:
        raise RuntimeError(data)
    return data["tenant_access_token"]


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v > 1e11:  # ms timestamp
            return datetime.fromtimestamp(v / 1000, tz=TZ).strftime("%Y/%m/%d")
        return str(v)
    if isinstance(v, list):
        parts = []
        for item in v:
            if isinstance(item, dict):
                parts.append(
                    item.get("text")
                    or item.get("name")
                    or item.get("email")
                    or ""
                )
            else:
                parts.append(_cell_text(item))
        return " / ".join(p for p in parts if p)
    if isinstance(v, dict):
        if "value" in v:
            return _cell_text(v.get("value"))
        return (
            v.get("text")
            or v.get("name")
            or _cell_text(v.get("value"))
            or ""
        )
    return str(v).strip()


def _list_all(token: str, table: str) -> list[dict]:
    """翻页 list records（比 search 全表更稳）。"""
    headers = {"Authorization": f"Bearer {token}"}
    items: list[dict] = []
    page_token = None
    while True:
        params: dict[str, Any] = {"page_size": 500}
        if page_token:
            params["page_token"] = page_token
        r = requests.get(
            f"https://open.feishu.cn/open-apis/bitable/v1/apps/{APP_TOKEN}/tables/{table}/records",
            headers=headers,
            params=params,
            timeout=90,
        )
        data = r.json()
        if data.get("code") != 0:
            raise RuntimeError(data)
        batch = data["data"].get("items") or []
        items.extend(batch)
        print(f"  …{table} 已拉取 {len(items)}", flush=True)
        if not data["data"].get("has_more"):
            break
        page_token = data["data"].get("page_token")
    return items


def _search_all(token: str, table: str, body: dict | None = None) -> list[dict]:
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    items: list[dict] = []
    page_token = None
    while True:
        payload = {"page_size": 500, **(body or {})}
        if page_token:
            payload["page_token"] = page_token
        r = requests.post(
            f"https://open.feishu.cn/open-apis/bitable/v1/apps/{APP_TOKEN}/tables/{table}/records/search",
            headers=headers,
            json=payload,
            timeout=90,
        )
        data = r.json()
        if data.get("code") != 0:
            raise RuntimeError(data)
        items.extend(data["data"].get("items") or [])
        print(f"  …{table} 已拉取 {len(items)}", flush=True)
        if not data["data"].get("has_more"):
            break
        page_token = data["data"].get("page_token")
    return items


def _parse_time(v: Any) -> datetime | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return datetime.fromtimestamp(v / 1000, tz=TZ)
    s = _cell_text(v)
    if not s:
        return None
    for fmt in ("%Y/%m/%d %H:%M", "%Y-%m-%dT%H:%M:%S%z", "%Y-%m-%d %H:%M:%S"):
        try:
            dt = datetime.strptime(s.replace("+08:00", "+0800"), fmt.replace("%z", "%z"))
            if dt.tzinfo is None:
                dt = dt.replace(tzinfo=TZ)
            return dt
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).astimezone(TZ)
    except ValueError:
        return None


def _fmt_date(v: Any) -> str:
    dt = _parse_time(v)
    return dt.strftime("%Y/%m/%d") if dt else _cell_text(v)


def _contact_info(fields: dict) -> str:
    parts = []
    for key, label in (
        (FIELD_EMAIL, "邮箱"),
        (FIELD_PHONE, "电话"),
        (FIELD_WECHAT, "微信"),
    ):
        val = _cell_text(fields.get(key))
        if val:
            parts.append(f"{label}:{val}")
    return "；".join(parts)


def _lead_source(fields: dict) -> str:
    ch = _cell_text(fields.get(FIELD_CHANNEL))
    sub = _cell_text(fields.get(FIELD_SUB))
    if ch and sub:
        return f"{ch}/{sub}"
    return ch or sub


def _customer_type(fields: dict) -> str:
    return (
        _cell_text(fields.get(FIELD_CUST_TYPE))
        or _cell_text(fields.get(FIELD_CUST_TYPE_US))
        or _cell_text(fields.get("Customer Grade"))
        or _cell_text(fields.get("Customer segmentation（客户分级）"))
    )


def _followup_value(fu: dict, key: str) -> str:
    fields = fu.get("fields") or {}
    if key == FU_TIME:
        return _fmt_date(fields.get(key))
    return _cell_text(fields.get(key))


def _related_lead_ids(related: Any) -> list[str]:
    """解析 Follow-up「Related Lead」关联 ID（兼容 list/search 两种返回形状）。"""
    link_ids: list[str] = []

    def walk(obj: Any) -> None:
        if not obj:
            return
        if isinstance(obj, str):
            link_ids.append(obj)
            return
        if isinstance(obj, dict):
            for key in ("link_record_ids", "record_ids"):
                vals = obj.get(key)
                if isinstance(vals, list):
                    link_ids.extend(str(x) for x in vals if x)
            for key in ("record_id", "id"):
                if obj.get(key):
                    link_ids.append(str(obj[key]))
            return
        if isinstance(obj, list):
            for item in obj:
                walk(item)

    walk(related)
    return list(dict.fromkeys(link_ids))


def build_rows(
    leads: list[dict],
    followups_by_lead: dict[str, list[dict]],
    salesperson: str | None,
) -> tuple[list[dict], int]:
    rows = []
    max_comms = 2
    for lead in leads:
        fields = lead.get("fields") or {}
        final = _cell_text(fields.get(FIELD_FINAL))
        if not final or final in ("未命中规则",):
            continue
        if salesperson and salesperson not in final and final not in salesperson:
            continue
        rid = lead["record_id"]
        fus = sorted(
            followups_by_lead.get(rid, []),
            key=lambda x: _parse_time((x.get("fields") or {}).get(FU_TIME))
            or datetime.min.replace(tzinfo=TZ),
        )
        max_comms = max(max_comms, len(fus), 2)
        rows.append(
            {
                "lead_id": _cell_text(fields.get(FIELD_LEAD_ID)),
                "company": _cell_text(fields.get(FIELD_UNIT)),
                "address": _cell_text(fields.get(FIELD_ADDRESS)),
                "contact": _cell_text(fields.get(FIELD_CUSTOMER)),
                "contact_info": _contact_info(fields),
                "source": _lead_source(fields),
                "reg_date": _fmt_date(fields.get(FIELD_ENTRY)),
                "cust_type": _customer_type(fields),
                "final": final,
                "followups": fus,
            }
        )
    rows.sort(key=lambda r: (r["final"], r["company"] or r["lead_id"]))
    return rows, max_comms


def write_memo(path: Path, rows: list[dict], max_comms: int, title_suffix: str = "") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "客户推进备忘"
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(wrap_text=True, horizontal="center", vertical="center")
    title = "客户推进备忘记录\nCustomer Follow-up Memo"
    if title_suffix:
        title = f"{title}\n({title_suffix})"

    total_cols = 8 + 10 * max_comms
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    ws.cell(1, 1).alignment = center

    # Row 2-3 headers
    for col, text in enumerate(CUSTOMER_HEADERS, 1):
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
        cell = ws.cell(2, col, text)
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.border = thin
        ws.cell(3, col).border = thin

    for i in range(max_comms):
        start = 9 + i * 10
        end = start + 9
        label = f"第{i + 1}次交流\nCommunication {i + 1}"
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=end)
        top = ws.cell(2, start, label)
        top.alignment = center
        top.font = Font(bold=True)
        for c in range(start, end + 1):
            ws.cell(2, c).border = thin
            h = COMM_HEADERS[c - start][0]
            cell = ws.cell(3, c, h)
            cell.alignment = center
            cell.font = Font(bold=True)
            cell.border = thin

    for idx, row in enumerate(rows, 1):
        r = 3 + idx
        values = [
            idx,
            row["company"],
            row["address"],
            row["contact"],
            row["contact_info"],
            row["source"],
            row["reg_date"],
            row["cust_type"],
        ]
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = center
            cell.border = thin
        fus = row["followups"]
        for i in range(max_comms):
            fu = fus[i] if i < len(fus) else None
            for j, (_label, key) in enumerate(COMM_HEADERS):
                c = 9 + i * 10 + j
                val = _followup_value(fu, key) if fu else ""
                cell = ws.cell(r, c, val)
                cell.alignment = center
                cell.border = thin

    widths = [6, 18, 20, 12, 22, 14, 12, 12] + [12] * (10 * max_comms)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 36

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="导出客户推进备忘录 Memo")
    parser.add_argument("--salesperson", help="最终分配业务员姓名（模糊匹配）；省略则全员分文件导出")
    parser.add_argument("--out", help="输出 xlsx 路径；省略则写入 exports/")
    parser.add_argument("--min-comms", type=int, default=2, help="最少交流列数，默认 2")
    args = parser.parse_args()

    token = _token()
    print("拉取线索…", flush=True)
    leads = _list_all(token, CASE_TABLE)
    if args.salesperson:
        leads = [
            it
            for it in leads
            if args.salesperson
            in _cell_text((it.get("fields") or {}).get(FIELD_FINAL))
        ]
    print(f"线索 {len(leads)} 条；拉取跟进…", flush=True)
    followups_raw = _list_all(token, FOLLOWUP_TABLE)
    by_lead: dict[str, list[dict]] = defaultdict(list)
    for fu in followups_raw:
        fields = fu.get("fields") or {}
        if fields.get(FU_DELETED) is True:
            continue
        for lid in _related_lead_ids(fields.get(FU_RELATED)):
            by_lead[lid].append(fu)
    print(f"跟进 {len(followups_raw)} 条，关联到 {len(by_lead)} 个线索", flush=True)

    rows, max_comms = build_rows(leads, by_lead, args.salesperson)
    max_comms = max(max_comms, args.min_comms)
    if not rows:
        print("无匹配客户，未生成文件")
        return 1

    out_dir = Path(__file__).resolve().parent.parent / "exports"
    if args.salesperson:
        path = Path(args.out) if args.out else out_dir / f"memo_{args.salesperson}_{datetime.now(TZ).strftime('%Y%m%d')}.xlsx"
        write_memo(path, rows, max_comms, title_suffix=args.salesperson)
        print(f"已导出 {len(rows)} 行 / {max_comms} 次交流列 → {path}")
        return 0

    # 全员：按最终分配业务员拆文件
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row["final"]].append(row)
    for name, group in sorted(grouped.items()):
        local_max = max(args.min_comms, max((len(r["followups"]) for r in group), default=0), 2)
        safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in name) or "unknown"
        path = out_dir / f"memo_{safe}_{datetime.now(TZ).strftime('%Y%m%d')}.xlsx"
        write_memo(path, group, local_max, title_suffix=name)
        print(f"  {name}: {len(group)} 行 → {path}")
    print(f"完成，共 {len(grouped)} 个业务员文件")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyError as e:
        print(f"缺少环境变量 {e}，请 source .env", file=sys.stderr)
        raise SystemExit(2)
